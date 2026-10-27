"""Focused, offline tests for extraction, chunking, retrieval, and RRF."""

import io
import unittest
from unittest.mock import patch

import numpy as np
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, DecodedStreamObject, NameObject

import rag_engine
from rag_engine import (
    DocumentChunk,
    SessionDocumentIndex,
    chunk_document,
    extract_text_from_file,
    prune_session_index,
    weighted_rrf,
)


class FakeEmbedding:
    """Small deterministic embedding substitute; no model download or API call."""

    def embed(self, texts):
        vectors = []
        for text in texts:
            lowered = text.lower()
            if "quantum" in lowered or "superposition" in lowered:
                vectors.append(np.array([1.0, 0.0, 0.0], dtype=np.float32))
            elif "gradient" in lowered or "optimization" in lowered:
                vectors.append(np.array([0.0, 1.0, 0.0], dtype=np.float32))
            else:
                vectors.append(np.array([0.0, 0.0, 1.0], dtype=np.float32))
        return vectors


def make_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("DOCX heading", level=1)
    doc.add_paragraph("DOCX extraction preserves this readable paragraph.")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def make_pdf_bytes() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = writer._add_object(
        DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
    )
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 200 Td (PDF extraction works.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class RagFeatureTests(unittest.TestCase):
    def test_extracts_txt_docx_and_pdf_with_metadata(self):
        txt = extract_text_from_file("notes.txt", b"Plain text extraction works.")
        self.assertEqual(txt, [("Plain text extraction works.", None)])

        docx = extract_text_from_file("notes.docx", make_docx_bytes())
        self.assertIn("DOCX extraction", docx[0][0])
        self.assertIsNone(docx[0][1])

        pdf = extract_text_from_file("notes.pdf", make_pdf_bytes())
        self.assertIn("PDF extraction works", pdf[0][0])
        self.assertEqual(pdf[0][1], 1)

    def test_invalid_and_empty_files_are_rejected(self):
        with self.assertRaises(ValueError):
            extract_text_from_file("empty.txt", b"   ")
        with self.assertRaises(ValueError):
            extract_text_from_file("fake.pdf", b"not a pdf")
        with self.assertRaises(ValueError):
            extract_text_from_file("fake.docx", b"not a zip")
        with self.assertRaises(ValueError):
            extract_text_from_file("binary.txt", b"hello\x00world")

    def test_chunking_preserves_overlap_and_metadata(self):
        sections = [("First paragraph about vectors.\n\nSecond paragraph about vectors.\n\nThird paragraph about vectors.", 4)]
        chunks = chunk_document("doc-1", "source.pdf", sections, chunk_size=55, overlap=12)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.doc_id == "doc-1" for chunk in chunks))
        self.assertTrue(all(chunk.page_number == 4 for chunk in chunks))
        self.assertTrue(any(set(chunks[i].text.split()) & set(chunks[i + 1].text.split()) for i in range(len(chunks) - 1)))

    def test_rrf_is_weighted_one_indexed_and_deduplicated(self):
        a = DocumentChunk("a", "d", "a.txt", None, "A", 0)
        b = DocumentChunk("b", "d", "b.txt", None, "B", 1)
        c = DocumentChunk("c", "d", "c.txt", None, "C", 2)
        fused = weighted_rrf([a, b], [b, c], semantic_weight=0.5, keyword_weight=0.5, rrf_k=60)
        self.assertEqual([chunk.chunk_id for chunk, _score in fused], ["b", "a", "c"])
        self.assertAlmostEqual(fused[0][1], 0.5 / 62 + 0.5 / 61)
        self.assertEqual(len({chunk.chunk_id for chunk, _score in fused}), 3)

    def test_hybrid_retrieval_is_semantic_keyword_and_limited_to_six(self):
        session_id = "offline-rag-feature-test"
        prune_session_index(session_id)
        try:
            with patch.object(rag_engine, "get_embedding_model", return_value=FakeEmbedding()):
                index = SessionDocumentIndex(session_id)
                for i in range(8):
                    index.add_document(f"quantum-{i}.txt", f"Quantum superposition concept {i}.".encode())
                index.add_document("optimization.txt", b"Gradient descent is an optimization algorithm.")
                results = index.hybrid_search("Explain quantum superposition", final_top_k=6)
                self.assertEqual(len(results), 6)
                self.assertTrue(all("quantum" in result.text.lower() for result in results))
                self.assertEqual(len({result.chunk_id for result in results}), len(results))
        finally:
            prune_session_index(session_id)


if __name__ == "__main__":
    unittest.main(verbosity=2)
